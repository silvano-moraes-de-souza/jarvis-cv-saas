# ================================================================
# JARVIS CV
# ARQUIVO: docx_generator.py
# DESCRIÇÃO: Pipeline DOCX Premium - Template ATS recruiter-friendly
# AUTOR: SILVANO MORAES DE SOUZA
# VERSÃO: 3.0.0
# ================================================================
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
import tempfile
import os
from typing import List, Optional

from app.utils.logger import logger
from app.config import settings

COLORS = {
    "primary": RGBColor(0x1a, 0x1a, 0x2e),
    "secondary": RGBColor(0x33, 0x33, 0x33),
    "accent": RGBColor(0x1e, 0x40, 0xaf),
    "light": RGBColor(0x64, 0x64, 0x64),
    "white": RGBColor(0xff, 0xff, 0xff),
}

FONT_MAIN = "Calibri"
FONT_ALT = "Arial"


class DocxGenerator:
    """Pipeline DOCX Premium - Template ATS recruiter-friendly sem mentira"""

    @staticmethod
    def _set_cell_shading(cell, color_hex: str):
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    @staticmethod
    def _add_horizontal_line(doc):
        p = doc.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        pBdr = parse_xml(
            f'<w:pBdr {nsdecls("w")}>'
            f'  <w:bottom w:val="single" w:sz="4" w:space="1" w:color="1e40af"/>'
            f'</w:pBdr>'
        )
        pPr.append(pBdr)
        return p

    @staticmethod
    async def generate_premium_resume(
        headline: str,
        summary: str,
        skills: List[str],
        experience_highlights: List[str],
        key_changes: List[dict] = None,
        contact_info: dict = None,
        original_name: str = "Curriculo_Otimizado",
    ) -> str:
        """
        Gera DOCX premium com layout ATS-friendly.
        Regras ATS aplicadas:
        - Sem tabelas para conteúdo (ATS não parseia)
        - Sem gráficos, imagens, colunas
        - Fonte padrão (Calibri/Arial)
        - Hierarquia clara com headings
        - Keywords naturais no texto
        - Verbos de poder no início dos bullets
        """
        doc = Document()

        # Configuração de página
        section = doc.sections[0]
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

        # Estilo base
        style = doc.styles["Normal"]
        font = style.font
        font.name = FONT_MAIN
        font.size = Pt(10.5)
        font.color.rgb = COLORS["secondary"]
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.line_spacing = 1.15

        # HEADLINE (Nome/Headline profissional)
        headline_para = doc.add_paragraph()
        headline_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = headline_para.add_run(headline)
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = COLORS["primary"]
        run.font.name = FONT_MAIN
        headline_para.paragraph_format.space_after = Pt(2)

        # Contato (se disponível)
        if contact_info:
            contact_parts = []
            if contact_info.get("email"):
                contact_parts.append(contact_info["email"])
            if contact_info.get("phone"):
                contact_parts.append(contact_info["phone"])
            if contact_info.get("linkedin"):
                contact_parts.append(contact_info["linkedin"])
            if contact_info.get("location"):
                contact_parts.append(contact_info["location"])

            if contact_parts:
                contact_para = doc.add_paragraph()
                run = contact_para.add_run(" | ".join(contact_parts))
                run.font.size = Pt(9)
                run.font.color.rgb = COLORS["accent"]
                run.font.name = FONT_MAIN
                contact_para.paragraph_format.space_after = Pt(6)

        # Linha divisória
        DocxGenerator._add_horizontal_line(doc)

        # RESUMO PROFISSIONAL
        section_title = doc.add_paragraph()
        run = section_title.add_run("RESUMO PROFISSIONAL")
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = COLORS["primary"]
        run.font.name = FONT_MAIN
        section_title.paragraph_format.space_before = Pt(10)
        section_title.paragraph_format.space_after = Pt(4)

        summary_para = doc.add_paragraph(summary)
        summary_para.paragraph_format.space_after = Pt(8)

        # COMPETÊNCIAS CHAVE
        if skills:
            section_title = doc.add_paragraph()
            run = section_title.add_run("COMPETÊNCIAS CHAVE")
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = COLORS["primary"]
            run.font.name = FONT_MAIN
            section_title.paragraph_format.space_before = Pt(10)
            section_title.paragraph_format.space_after = Pt(4)

            skills_text = " | ".join(skills[:15])
            skills_para = doc.add_paragraph()
            run = skills_para.add_run(skills_text)
            run.font.size = Pt(10)
            run.font.color.rgb = COLORS["accent"]
            run.font.name = FONT_MAIN
            skills_para.paragraph_format.space_after = Pt(8)

        # DESTAQUES DE EXPERIÊNCIA
        if experience_highlights:
            DocxGenerator._add_horizontal_line(doc)

            section_title = doc.add_paragraph()
            run = section_title.add_run("DESTAQUES DE EXPERIÊNCIA")
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = COLORS["primary"]
            run.font.name = FONT_MAIN
            section_title.paragraph_format.space_before = Pt(10)
            section_title.paragraph_format.space_after = Pt(4)

            for highlight in experience_highlights:
                bullet = doc.add_paragraph(style="List Bullet")
                run = bullet.add_run(highlight)
                run.font.size = Pt(10)
                run.font.color.rgb = COLORS["secondary"]
                run.font.name = FONT_MAIN
                bullet.paragraph_format.space_after = Pt(2)

        # MUDANÇAS CHAVE (se disponível)
        if key_changes:
            doc.add_page_break()

            section_title = doc.add_paragraph()
            run = section_title.add_run("RELATÓRIO DE OTIMIZAÇÃO JARVIS CV")
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = COLORS["primary"]
            run.font.name = FONT_MAIN
            section_title.paragraph_format.space_after = Pt(8)

            DocxGenerator._add_horizontal_line(doc)

            for change in key_changes:
                change_para = doc.add_paragraph()
                section_name = change.get("section", "")
                run = change_para.add_run(f"▸ {section_name}")
                run.font.bold = True
                run.font.size = Pt(10)
                run.font.color.rgb = COLORS["accent"]

                if change.get("before"):
                    before_para = doc.add_paragraph()
                    run = before_para.add_run(f"  Antes: {change['before']}")
                    run.font.size = Pt(9)
                    run.font.color.rgb = COLORS["light"]
                    run.font.italic = True

                if change.get("after"):
                    after_para = doc.add_paragraph()
                    run = after_para.add_run(f"  Depois: {change['after']}")
                    run.font.size = Pt(9)
                    run.font.color.rgb = COLORS["secondary"]
                    run.font.bold = True

                if change.get("why"):
                    why_para = doc.add_paragraph()
                    run = why_para.add_run(f"  Por quê: {change['why']}")
                    run.font.size = Pt(9)
                    run.font.color.rgb = COLORS["light"]
                    why_para.paragraph_format.space_after = Pt(6)

        # Footer
        footer_para = doc.add_paragraph()
        footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        contact_email = os.getenv("JARVIS_CONTACT_EMAIL", "contato@jarviscv.com")
        run = footer_para.add_run(f"Otimizado por JARVIS CV | {contact_email}")
        run.font.size = Pt(8)
        run.font.color.rgb = COLORS["light"]
        run.font.italic = True

        # Salva
        temp_dir = tempfile.mkdtemp()
        filename = f"{original_name}_JARVIS_CV.docx"
        filepath = os.path.join(temp_dir, filename)
        doc.save(filepath)

        logger.log(f"[DOCX] Gerado: {filepath}")
        return filepath


docx_generator = DocxGenerator()
